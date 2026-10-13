#!/usr/bin/env python3
"""Lecture slides (PDF or PPTX) -> slide text, slide thumbnails and cropped figures. Offline.

    python3 slides.py <lecture.pdf|lecture.pptx> [--name "Original name.pptx"] [--debug DIR]
    python3 slides.py --check      what's installed, as JSON

Prints one JSON object on stdout. Progress goes to stderr as "PROGRESS <stage> <i> <n>" lines,
which the notebook's /api/extract route turns into the progress bar.

A .pptx is first printed to PDF by LibreOffice (headless), so both kinds go through the same
figure finder; python-pptx adds what a PDF loses (speaker notes, the real slide titles).

Slides have no "Figure 3:" captions to anchor on (that's how PaperLamp finds figures in papers),
so a figure here is found by what it's made of:
  - pictures placed on the slide, minus the deck's template: anything repeated on many slides
    (logos, backgrounds, header bands), edge-to-edge decorative strips, and tiny icons;
  - vector drawings (flowcharts, trees, tables, charts), clustered into one region per diagram,
    kept only when they look like a diagram (connectors, grid lines, many shapes with short
    labels) rather than a styled text box;
  - then neighbouring pieces are merged, the short labels around a diagram are pulled in, and the
    region is rendered from the page, so crops, masks and recolouring look exactly as on screen.
The notebook then shows every crop to Gemma (vision), which drops the false positives and writes
the caption, so this file errs on the side of finding too much rather than too little.
"""
import base64
import collections
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

try:
    import fitz  # PyMuPDF
    from PIL import Image
except ImportError as missing:  # reported by --check and as a readable error, not a traceback
    fitz = None
    MISSING = str(missing)

THUMB_W = 560          # slide thumbnails, px
FIG_MAX = 1400         # longest side of a figure crop, px
MAX_FIGS_PER_SLIDE = 3
BULLET = re.compile(r"^[•▪●◦‣⁃∙➢➔►▸✓✔\-–—\*]\s*")
PRIVATE = re.compile(r"[-]")  # Wingdings/Symbol bullets come out of PowerPoint as private-use glyphs
SKIP_TITLES = re.compile(r"^(thank\s*you|thanks|questions?\??|q\s*&\s*a|any questions\??|the end)\b", re.I)


def progress(stage, i, n):
    print(f"PROGRESS {stage} {i} {n}", file=sys.stderr, flush=True)


# ── PPTX → PDF ──────────────────────────────────────────────────────────────
def soffice():
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    mac = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    return mac if os.path.exists(mac) else None


def pptx_to_pdf(path, work):
    """Prints the deck to PDF with LibreOffice. A private profile, so it works even while the
    user has LibreOffice open."""
    exe = soffice()
    if not exe:
        return None
    profile = "file://" + os.path.join(work, "lo-profile")
    subprocess.run([exe, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf", "--outdir", work, path],
                   capture_output=True, timeout=600)
    pdf = os.path.join(work, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    return pdf if os.path.exists(pdf) else None


def pptx_info(path):
    """Titles, speaker notes and the deck title from the .pptx itself (visible slides only, in
    order, so they line up with the pages LibreOffice printed)."""
    try:
        from pptx import Presentation
    except ImportError:
        return None
    prs = Presentation(path)
    slides = []
    for slide in prs.slides:
        if slide._element.get("show") == "0":  # hidden slides aren't printed
            continue
        title = ""
        try:
            if slide.shapes.title is not None:
                title = slide.shapes.title.text_frame.text
        except Exception:
            pass
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text
        except Exception:
            pass
        slides.append(dict(title=clean(title), notes=clean(notes, keep_lines=True)))
    return dict(title=clean(prs.core_properties.title or ""), slides=slides,
                width=prs.slide_width, height=prs.slide_height)


# ── text ────────────────────────────────────────────────────────────────────
def clean(s, keep_lines=False):
    s = PRIVATE.sub("•", s or "").replace(" ", " ").replace("\u000b", "\n")
    if keep_lines:
        return "\n".join(re.sub(r"[ \t]+", " ", ln).strip() for ln in s.splitlines() if ln.strip())
    return re.sub(r"\s+", " ", s).strip()


def norm(s):
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", s.lower())).strip()


def page_lines(page):
    """Every line of text with its box, size and boldness."""
    out = []
    for b in page.get_text("dict")["blocks"]:
        if b.get("type") != 0:
            continue
        for ln in b["lines"]:
            spans = [s for s in ln["spans"] if s["text"].strip()]
            if not spans:
                continue
            text = clean("".join(s["text"] for s in ln["spans"]))
            if not text:
                continue
            out.append(dict(text=text, box=fitz.Rect(ln["bbox"]), size=round(max(s["size"] for s in spans), 1),
                            bold=any(s["flags"] & 16 or "bold" in s["font"].lower() for s in spans), block=b["number"]))
    return out


def template_lines(all_lines, n):
    """Lines repeated on many slides (course code, faculty name, "Prof. X") are the template."""
    seen = collections.Counter()
    for lines in all_lines:
        seen.update({norm(l["text"]) for l in lines})
    cut = max(3, 0.4 * n)
    return {t for t, c in seen.items() if c >= cut}


def split_title(lines, page, template):
    """The slide title: the biggest text in the top part of the slide."""
    H = page.rect.height
    body = [l for l in lines if norm(l["text"]) not in template and not re.fullmatch(r"\d{1,3}", l["text"])]
    if not body:
        return "", []
    # a title starts like one (capital, digit or quote), isn't a bullet, and isn't a whole sentence
    top = [l for l in body if l["box"].y0 < H * 0.4 and not BULLET.match(l["text"]) and len(l["text"].split()) <= 16
           and re.match(r"[\"'“‘(]?[A-Z0-9]", l["text"])]
    title_lines = []
    if top:
        big = max(l["size"] for l in top)
        rest = [l["size"] for l in body if l["size"] < big]
        median = sorted(rest)[len(rest) // 2] if rest else big
        if big >= median * 1.15:
            first = min((l for l in top if l["size"] >= big - 0.6), key=lambda l: (l["box"].y0, l["box"].x0))
            # a title wrapped onto two or three lines: same size, stacked right under the first
            title_lines = [first]
            for l in sorted(body, key=lambda l: l["box"].y0):
                last = title_lines[-1]
                if l is not first and l not in title_lines and abs(l["size"] - big) < 0.7 and 0 < l["box"].y0 - last["box"].y0 < big * 1.6:
                    title_lines.append(l)
    if not title_lines and top:
        # no bigger text: the first line counts only if it stands alone, with a gap under it
        first = min(body, key=lambda l: (l["box"].y0, l["box"].x0))
        below = [l for l in body if l is not first and l["box"].y0 >= first["box"].y1 - 1]
        gap = min((l["box"].y0 - first["box"].y1 for l in below), default=99)
        if first in top and first["box"].y0 < H * 0.25 and len(first["text"].split()) <= 10 and gap > first["size"] * 0.8:
            title_lines = [first]
    title = clean(" ".join(l["text"] for l in sorted(title_lines, key=lambda l: (l["box"].y0, l["box"].x0))))
    return title, [l for l in body if l not in title_lines]


def body_text(lines):
    """The slide's text in reading order, with bullets and indentation kept as plain text."""
    if not lines:
        return ""
    left = min(l["box"].x0 for l in lines)
    rows = sorted(lines, key=lambda l: (round(l["box"].y0 / 6), l["box"].x0))
    out, prev = [], None
    for l in rows:
        text = BULLET.sub("", l["text"]).strip()
        if not text or text == "•":
            continue
        bullet = bool(BULLET.match(l["text"])) or l["text"].startswith("•")
        indent = "  " * min(3, int((l["box"].x0 - left) // 28))
        # a wrapped line: close under the previous one, roughly the same indent, no bullet of its own
        same_para = (prev is not None and not bullet and abs(l["size"] - prev["size"]) < 1.5
                     and -2 < l["box"].y0 - prev["box"].y1 < l["size"] * 0.9
                     and (abs(l["box"].x0 - prev["box"].x0) < 40 or text[:1].islower()))
        if same_para and out:
            out[-1] += " " + text
        else:
            out.append(f"{indent}{'- ' if bullet else ''}{text}")
        prev = l
    return "\n".join(out)


# ── figures ─────────────────────────────────────────────────────────────────
def image_key(info):
    return info.get("digest") or info.get("xref") or (info["width"], info["height"], info.get("size"))


def drawing_key(d):
    r = d["rect"]
    col = lambda c: tuple(round(x, 2) for x in c) if c else None
    return (round(r.x0), round(r.y0), round(r.x1), round(r.y1), col(d.get("fill")), col(d.get("color")))


def image_regions(page, infos, img_count, n):
    W, H = page.rect.width, page.rect.height
    area = W * H
    big, small = [], []
    for info in infos:
        r = fitz.Rect(info["bbox"]) & page.rect
        if r.is_empty or r.width < 4 or r.height < 4:
            continue
        if img_count[image_key(info)] >= max(3, 0.3 * n):
            continue  # logo, background or header band repeated through the deck
        a = r.width * r.height
        band = (r.width > 0.9 * W and r.height < 0.4 * H) or (r.height > 0.9 * H and r.width < 0.5 * W)
        if band:
            continue  # decorative strip along an edge
        if a > 0.72 * area:
            continue  # full-bleed background; a slide that is one big picture is kept below
        if r.width < 0.07 * W or r.height < 0.07 * H or a < 0.012 * area:
            small.append(r)  # icons: only kept as part of a diagram
            continue
        big.append(r)
    return big, small


def vector_regions(page, drawings, draw_count, n, lines):
    """Clusters of vector shapes that look like a diagram, table or chart."""
    W, H = page.rect.width, page.rect.height
    area = W * H
    keep = []
    for d in drawings:
        r = d["rect"]
        if draw_count[drawing_key(d)] >= max(3, 0.3 * n):
            continue
        if r.width * r.height > 0.6 * area or (r.width > 0.95 * W and r.height > 0.3 * H):
            continue
        if d.get("fill") is None and d.get("color") is None:
            continue
        if r.width < 0.5 and r.height < 0.5:
            continue
        keep.append(d)
    if not keep:
        return []

    # union-find over shapes that touch or nearly touch
    tol = 9
    parent = list(range(len(keep)))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    boxes = [fitz.Rect(d["rect"]) + (-tol, -tol, tol, tol) for d in keep]
    for i in range(len(keep)):
        for j in range(i + 1, len(keep)):
            if boxes[i].intersects(boxes[j]):
                parent[find(i)] = find(j)
    groups = collections.defaultdict(list)
    for i, d in enumerate(keep):
        groups[find(i)].append(d)

    out = []
    for members in groups.values():
        r = fitz.Rect()
        for d in members:
            r |= d["rect"]
        a = r.width * r.height
        if a < 0.035 * area or r.width < 0.12 * W or r.height < 0.08 * H:
            continue
        closed = [d for d in members if (d.get("closePath") or d.get("fill") is not None) and d["rect"].width > 8 and d["rect"].height > 8]
        open_strokes = [d for d in members if d.get("fill") is None and not d.get("closePath")]
        connectors = sum(1 for d in open_strokes if links(d, closed))
        # table rules: stroked lines or hairline-thin filled bars (not the edges of filled cards)
        rules = [(it[1], it[2]) for d in open_strokes for it in d["items"] if it[0] == "l"]
        rules += [(d["rect"].tl, d["rect"].br) for d in members if d.get("fill") is not None and min(d["rect"].width, d["rect"].height) < 2]
        hlines = sum(1 for p, q in rules if abs(p.y - q.y) < 2 and abs(p.x - q.x) > r.width * 0.3)
        vlines = sum(1 for p, q in rules if abs(p.x - q.x) < 2 and abs(p.y - q.y) > r.height * 0.3)
        free_curves = sum(1 for d in open_strokes for it in d["items"] if it[0] == "c")
        inside = [l for l in lines if r.contains(l["box"].tl + (1, 1))]
        words = [len(l["text"].split()) for l in inside]
        avg_words = sum(words) / len(words) if words else 0
        table = hlines >= 3 and vlines >= 2 or hlines >= 5
        # boxes joined by lines or arrows (flowcharts, trees, cycles), or free-hand curves (charts,
        # timelines); a grid of cards with no lines between them is layout, not a figure
        diagram = connectors >= 2 or (connectors >= 1 and len(closed) >= 3) or free_curves >= 6
        wordy = avg_words > 7.5 and not table
        if (table or diagram) and not wordy:
            out.append(r)
    return out


def links(d, shapes):
    """Whether an open line runs from one shape to a different one (a connector or an arrow),
    rather than being a divider or an underline inside a single card."""
    pts = [p for it in d["items"] for p in it[1:] if hasattr(p, "x")]
    if len(pts) < 2:
        return False
    a, b = pts[0], pts[-1]
    if abs(a.x - b.x) + abs(a.y - b.y) < 12:
        return False

    def touching(p):
        return {i for i, s in enumerate(shapes) if (s["rect"] + (-7, -7, 7, 7)).contains(p)
                and not (s["rect"] + (5, 5, -5, -5)).contains(p)}

    ta, tb = touching(a), touching(b)
    if ta and tb:
        return bool(ta ^ tb) or not (ta & tb)
    # an arrow from a shape out to a label or another line
    return bool(ta or tb) and abs(a.x - b.x) + abs(a.y - b.y) > 30 and not any(
        s["rect"].contains(a) and s["rect"].contains(b) for s in shapes)


def area(r):
    return 0 if r.is_empty else r.width * r.height


def merge(rects, gap=14):
    rects = [fitz.Rect(r) for r in rects]
    changed = True
    while changed:
        changed = False
        for i in range(len(rects)):
            for j in range(i + 1, len(rects)):
                if (rects[i] + (-gap, -gap, gap, gap)).intersects(rects[j]):
                    rects[i] |= rects[j]
                    rects.pop(j)
                    changed = True
                    break
            if changed:
                break
    return rects


def with_labels(r, lines, title_box):
    """Pulls in short labels that sit just outside a diagram ("Search", "Answer" beside an arrow)."""
    grown = fitz.Rect(r)
    for _ in range(2):
        for l in lines:
            if title_box and l["box"].intersects(title_box):
                continue
            if len(l["text"].split()) > 6:
                continue
            near = grown + (-16, -16, 16, 16)
            if near.intersects(l["box"]) and not grown.contains(l["box"]):
                grown |= l["box"]
    return grown


def render(page, r, photo):
    scale = max(2.0, min(4.0, FIG_MAX / max(r.width, r.height)))
    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), clip=r, alpha=False)
    im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    buf = io.BytesIO()
    if photo:
        im.save(buf, "JPEG", quality=84, optimize=True)
        mime = "image/jpeg"
    else:
        im.quantize(colors=128, method=2).save(buf, "PNG", optimize=True)  # 2 = fast octree; line art needs few colours
        mime = "image/png"
    return f"data:{mime};base64," + base64.b64encode(buf.getvalue()).decode(), im.width, im.height


def thumb(page):
    scale = THUMB_W / page.rect.width
    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
    im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=66, optimize=True)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()


# ── the whole deck ──────────────────────────────────────────────────────────
def extract(pdf_path, name, kind, extra=None):
    doc = fitz.open(pdf_path)
    n = doc.page_count
    W, H = doc[0].rect.width, doc[0].rect.height

    # pass 1: what repeats through the deck (the template)
    all_lines, all_infos, all_draws = [], [], []
    img_count, draw_count = collections.Counter(), collections.Counter()
    for i, page in enumerate(doc):
        progress("read", i + 1, n)
        all_lines.append(page_lines(page))
        infos = page.get_image_info(hashes=True, xrefs=True)
        all_infos.append(infos)
        img_count.update({image_key(x) for x in infos})
        draws = page.get_drawings()
        all_draws.append(draws)
        draw_count.update({drawing_key(d) for d in draws})
    template = template_lines(all_lines, n)

    # pass 2: text, thumbnails, figures
    slides, figures = [], []
    pp = (extra or {}).get("slides") or []
    for i, page in enumerate(doc):
        progress("figures", i + 1, n)
        title, body = split_title(all_lines[i], page, template)
        info = pp[i] if i < len(pp) else {}
        if info.get("title"):
            title = info["title"]
        text = body_text(body)
        words = len(text.split())
        title_box = None
        for l in all_lines[i]:
            if title and l["text"] and l["text"] in title:
                title_box = l["box"] if title_box is None else title_box | l["box"]

        big, small = image_regions(page, all_infos[i], img_count, n)
        vec = vector_regions(page, all_draws[i], draw_count, n, body)
        regions = merge(big + vec)
        # icons only count when they sit in or next to a diagram
        regions = merge(regions + [s for s in small if any((s + (-20, -20, 20, 20)).intersects(r) for r in regions)])
        regions = [with_labels(r, body, title_box) for r in regions]
        if title_box is not None:
            regions = [fitz.Rect(r.x0, max(r.y0, title_box.y1 + 2), r.x1, r.y1) if r.intersects(title_box) and r.y1 > title_box.y1 + 40 else r
                       for r in regions]
        regions = merge(regions, gap=4)
        regions = [r for r in regions if r.width * r.height >= 0.03 * W * H]
        # a slide that is basically one big picture: keep the picture, whatever its size
        if not regions and words < 12:
            for x in all_infos[i]:
                r = fitz.Rect(x["bbox"]) & page.rect
                if img_count[image_key(x)] < max(3, 0.3 * n) and r.width * r.height > 0.25 * W * H:
                    regions.append(r)
                    break
        regions.sort(key=lambda r: -(r.width * r.height))
        regions = sorted(regions[:MAX_FIGS_PER_SLIDE], key=lambda r: (r.y0, r.x0))

        for r in regions:
            pad = fitz.Rect(r.x0 - 6, r.y0 - 6, r.x1 + 6, r.y1 + 6) & page.rect
            photo = any(area(fitz.Rect(x["bbox"]) & pad) > 0.6 * area(pad) for x in all_infos[i]
                        if img_count[image_key(x)] < max(3, 0.3 * n))
            url, w, h = render(page, pad, photo)
            labels = clean(" · ".join(l["text"] for l in body if pad.contains(l["box"].tl + (1, 1))))[:400]
            figures.append(dict(id=f"f{len(figures) + 1}", slide=i + 1, src=url, w=w, h=h,
                                box=[round(pad.x0 / W, 4), round(pad.y0 / H, 4), round(pad.x1 / W, 4), round(pad.y1 / H, 4)],
                                kind="picture" if photo else "drawing", labels=labels))

        skip = bool(SKIP_TITLES.match(title or text[:40])) and words < 40
        if skip:  # the "Thank you" slide's social links aren't figures
            figures = [f for f in figures if f["slide"] != i + 1]
        slides.append(dict(n=i + 1, title=title, text=text, notes=info.get("notes", ""), words=words, skip=skip,
                           figures=[f["id"] for f in figures if f["slide"] == i + 1], thumb=thumb(page)))

    meta_title = clean(doc.metadata.get("title") or "")
    title = (extra or {}).get("title") or ""
    if not useful_title(title, name):
        title = meta_title if useful_title(meta_title, name) else ""
    if not title and slides:
        title = slides[0]["title"] or ""
    return dict(source=dict(name=name, kind=kind, pages=n, width=W, height=H), title=title or os.path.splitext(name)[0],
                slides=slides, figures=figures)


def useful_title(t, name):
    if not t or len(t) < 4:
        return False
    low = t.lower()
    return not (low.startswith("powerpoint") or low.startswith("microsoft") or low.startswith("slide")
                or low.endswith((".ppt", ".pptx", ".pdf", ".key")) or low == os.path.splitext(name)[0].lower())


def debug_sheet(result, folder):
    os.makedirs(folder, exist_ok=True)
    for f in result["figures"]:
        data = base64.b64decode(f["src"].split(",", 1)[1])
        ext = "jpg" if "jpeg" in f["src"][:30] else "png"
        with open(os.path.join(folder, f"s{f['slide']:03d}-{f['id']}.{ext}"), "wb") as fh:
            fh.write(data)
    slim = dict(result, figures=[dict(f, src="…") for f in result["figures"]],
                slides=[dict(s, thumb="…") for s in result["slides"]])
    with open(os.path.join(folder, "result.json"), "w") as fh:
        json.dump(slim, fh, indent=1, ensure_ascii=False)


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__, file=sys.stderr)
        sys.exit(2)
    if args[0] == "--check":
        try:
            import pptx  # noqa: F401
            has_pptx = True
        except ImportError:
            has_pptx = False
        print(json.dumps(dict(pymupdf=fitz is not None, pptx=has_pptx, libreoffice=bool(soffice()),
                              problem=None if fitz else f"{MISSING}. Run: pip3 install pymupdf pillow python-pptx")))
        return
    if fitz is None:
        print(json.dumps(dict(error=f"The slide reader needs PyMuPDF ({MISSING}). Run: pip3 install pymupdf pillow python-pptx")))
        sys.exit(1)
    path = args[0]
    debug = args[args.index("--debug") + 1] if "--debug" in args else None
    name = args[args.index("--name") + 1] if "--name" in args else os.path.basename(path)
    ext = os.path.splitext(name)[1].lower()
    work = tempfile.mkdtemp(prefix="studyink-")
    try:
        if ext in (".pptx", ".ppt", ".odp", ".key"):
            progress("convert", 0, 1)
            extra = pptx_info(path) if ext == ".pptx" else None
            pdf = pptx_to_pdf(path, work)
            if not pdf:
                print(json.dumps(dict(error="Couldn't print the slides. Install LibreOffice (brew install --cask libreoffice), or export the deck as PDF.")))
                sys.exit(1)
            progress("convert", 1, 1)
            result = extract(pdf, name, "pptx", extra)
        elif ext == ".pdf":
            result = extract(path, name, "pdf")
        else:
            print(json.dumps(dict(error=f"{ext or 'that file'} isn't a slide deck I can read. Use .pptx or .pdf.")))
            sys.exit(1)
        if debug:
            debug_sheet(result, debug)
        json.dump(result, sys.stdout, ensure_ascii=False)
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    main()
